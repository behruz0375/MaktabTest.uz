"""Fayldan (Excel/CSV/Word/PDF/Rasm) test savollarini aqlli o'qish.

Har qanday formatni tushunadi:
- Excel/CSV: ustunlar tartibi va nomlari ixtiyoriy (aqlli aniqlaydi)
- Word/PDF/Matn: raqamli/raqamsiz savollar, inline variantlar, turli belgilar
- Rasm: OCR matnini text parser bilan tahlil qiladi
"""
import io
import re
import csv


def _clean(s):
    return (str(s) if s is not None else "").strip()


def _detect_letter(correct_raw, a, b, c, d):
    raw = _clean(correct_raw)
    if not raw:
        return ""
    m = re.fullmatch(r"[ABCDabcd]\s*[\.\)\-:]?", raw)
    if m:
        return raw[0].upper()
    m2 = re.search(
        r"(?:javob|to'?g'?ri|answer|correct|key|kalit)\D{0,15}?\b([ABCDabcd])\b",
        raw, re.I,
    )
    if m2:
        return m2.group(1).upper()
    m3 = re.fullmatch(r"\s*([ABCDabcd])\s*", raw)
    if m3:
        return m3.group(1).upper()
    norm = raw.strip().lower()
    for letter, val in (("A", a), ("B", b), ("C", c), ("D", d)):
        if val and val.strip().lower() == norm:
            return letter
    return ""


def _q(text, a, b, c, d, correct):
    text, a, b, c, d = _clean(text), _clean(a), _clean(b), _clean(c), _clean(d)
    letter = _detect_letter(correct, a, b, c, d)
    if text and a and b and c and d and letter in ("A", "B", "C", "D"):
        return {"text": text, "option_a": a, "option_b": b, "option_c": c,
                "option_d": d, "correct_option": letter}
    return None


# =========================================================================
#  AQLLI USTUN ANIQLASH (Excel/CSV/Word jadvali)
# =========================================================================

def _classify_header(cell):
    """Ustun sarlavhasini rol bo'yicha aniqlash: question/A/B/C/D/answer/number."""
    h = _clean(cell)
    if not h:
        return None
    hl = h.lower()

    # Raqam ustuni
    if h in ("№", "#") or hl in ("t/r", "no", "raqam", "number", "tartib",
                                       "номер", "n", "id"):
        return "number"

    # Javob ustuni
    for kw in ("to‘g‘ri", "to'g'ri", "togri", "javob", "answer", "correct",
               "kalit", "ответ", "правильн",
               "key", "kalitlar", "toʻgʻri"):
        if kw in hl:
            return "answer"

    # Savol ustuni
    for kw in ("savol", "question", "вопрос", "matn", "text",
               "test", "savollar"):
        if kw in hl:
            return "question"

    # Variant ustunlari — bitta harf yoki "X variant" shakli
    if len(h) <= 3:
        ch = h.strip(".)- \t")
        if ch in ("A", "a", "А", "а", "1"):
            return "A"
        if ch in ("B", "b", "Б", "б", "В", "в", "2"):
            return "B"
        if ch in ("C", "c", "С", "с", "3"):
            return "C"
        if ch in ("D", "d", "Д", "д", "Г", "г", "4"):
            return "D"

    m = re.fullmatch(
        r"(?:variant|вариант|javob|option)\s*([ABCDabcd1234АаБбВвГгСсДд])"
        r"|([ABCDabcd1234АаБбВвГгСсДд])\s*(?:variant|вариант|javob|option)",
        h, re.I,
    )
    if m:
        ch = (m.group(1) or m.group(2)).upper()
        if ch in ("A", "А", "1"):
            return "A"
        if ch in ("B", "Б", "В", "2"):
            return "B"
        if ch in ("C", "С", "3"):
            return "C"
        if ch in ("D", "Д", "Г", "4"):
            return "D"

    return None


def _smart_columns(all_rows):
    """Sarlavha va ma'lumotlar tahlili orqali ustun rollarini aniqlash."""
    if not all_rows:
        return None, []

    # 1) Sarlavha qatorini topish (birinchi 5 qatordan birini tekshirish)
    header_idx = -1
    col_map = {}

    for idx in range(min(5, len(all_rows))):
        row = [_clean(c) for c in all_rows[idx]]
        if not any(row):
            continue
        assigned = {}
        for ci, cell in enumerate(row):
            role = _classify_header(cell)
            if role and role not in assigned:
                assigned[role] = ci
        n_opts = sum(1 for k in "ABCD" if k in assigned)
        if ("question" in assigned and n_opts >= 2) or \
           ("answer" in assigned and n_opts >= 2) or \
           n_opts >= 3:
            header_idx = idx
            col_map = assigned
            break

    # Data qatorlari
    if header_idx >= 0:
        data_rows = all_rows[header_idx + 1:]
    else:
        data_rows = all_rows
    data_rows = [r for r in data_rows if any(_clean(c) for c in r)]

    if not data_rows:
        return None, []

    ncols = max(len(r) for r in data_rows)
    used = set(col_map.values())

    # 2) Javob ustunini content'dan aniqlash (yagona harf A/B/C/D)
    if "answer" not in col_map:
        best_c, best_n = -1, 0
        for c in range(ncols):
            if c in used:
                continue
            n = sum(1 for r in data_rows
                    if c < len(r) and re.fullmatch(r"\s*[ABCDabcd]\s*[\.\)\-:]?\s*", _clean(r[c])))
            if n > best_n:
                best_n, best_c = n, c
        if best_c >= 0 and best_n >= len(data_rows) * 0.4:
            col_map["answer"] = best_c
            used.add(best_c)

    # 3) Raqam ustunini aniqlash (ketma-ket sonlar)
    if "number" not in col_map:
        for c in range(ncols):
            if c in used:
                continue
            nums = []
            for r in data_rows[:30]:
                v = _clean(r[c] if c < len(r) else "")
                if re.fullmatch(r"\d+", v):
                    nums.append(int(v))
                elif v:
                    break
            if len(nums) >= min(len(data_rows), 3) and nums == sorted(nums):
                col_map["number"] = c
                used.add(c)
                break

    # 4) Savol ustunini aniqlash (eng uzun o'rtacha matn)
    if "question" not in col_map:
        best_c, best_len = -1, 0
        for c in range(ncols):
            if c in used:
                continue
            avg = sum(len(_clean(r[c]) if c < len(r) else "") for r in data_rows) / max(len(data_rows), 1)
            if avg > best_len:
                best_len, best_c = avg, c
        if best_c >= 0 and best_len > 3:
            col_map["question"] = best_c
            used.add(best_c)

    # 5) Qolgan ustunlarni A, B, C, D ga tartib bo'yicha biriktirish
    remaining = sorted(c for c in range(ncols) if c not in used)
    for opt in ("A", "B", "C", "D"):
        if opt not in col_map and remaining:
            col_map[opt] = remaining.pop(0)

    if not all(k in col_map for k in ("question", "A", "B", "C", "D")):
        return None, []

    return col_map, data_rows


def _parse_table_rows(all_rows):
    """Jadval qatorlaridan savollarni chiqarish (aqlli ustun aniqlash bilan)."""
    # 1-usul: Aqlli ustun aniqlash
    col_map, data_rows = _smart_columns(all_rows)
    if col_map and data_rows:
        out = []
        max_ci = max(col_map.values()) + 1
        for row in data_rows:
            cells = list(row) + [None] * max_ci
            q = _q(cells[col_map["question"]],
                   cells[col_map["A"]], cells[col_map["B"]],
                   cells[col_map["C"]], cells[col_map["D"]],
                   cells[col_map.get("answer", -1)] if "answer" in col_map else "")
            if q:
                out.append(q)
        if out:
            return out

    # 2-usul: Eski qat'iy tartib (0-5 ustunlar)
    out = []
    for row in all_rows:
        cells = list(row) + [None] * 6
        joined = " ".join(_clean(c).lower() for c in cells[:6])
        if "savol" in joined or "javob" in joined or "question" in joined or "answer" in joined:
            continue
        q = _q(cells[0], cells[1], cells[2], cells[3], cells[4], cells[5])
        if q:
            out.append(q)
    if out:
        return out

    # 3-usul: Katakchalar matnini birlashtirish va text parser bilan o'qish
    lines = []
    for row in all_rows:
        for cell in row:
            text = _clean(cell)
            if text:
                lines.extend(text.splitlines())
    return parse_text_lines(lines)


def parse_excel(data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    all_rows = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            all_rows.append(list(row))
    return _parse_table_rows(all_rows)


def parse_csv(data):
    text = data.decode("utf-8-sig", errors="replace")
    delim = ";" if text.count(";") > text.count(",") else ","
    all_rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    return _parse_table_rows(all_rows)


def parse_docx(data):
    import docx
    doc = docx.Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs]
    table_rows = []
    for t in doc.tables:
        for r in t.rows:
            table_rows.append([c.text for c in r.cells])
    if table_rows:
        result = _parse_table_rows(table_rows)
        if result:
            return result
    return parse_text_lines(lines)


def parse_pdf(data):
    import pypdf
    reader = pypdf.PdfReader(io.BytesIO(data))
    lines = []
    for page in reader.pages:
        txt = page.extract_text() or ""
        lines.extend(txt.splitlines())
    return parse_text_lines(lines)


# =========================================================================
#  MATN PARSER (OCR va fayllar uchun)
# =========================================================================

OPT_RE = re.compile(r"^\s*([ABCDabcd])\s*(?:[\).\-:\]]\s*|\s+)(.+)$")
ANS_RE = re.compile(
    r"^\s*(javob(?:i)?|to'?g'?ri\s*javob(?:i)?|to'?g'?risi|to'?g'?ri|answer(?:\s*key)?|correct(?:\s*answer)?|key|kalit(?:\s*javob)?)\s*[:\-]?\s*([ABCDabcd])\b",
    re.I,
)
BARE_ANS_RE = re.compile(r"^\s*([ABCDabcd])\s*[\.\)]?\s*$")
Q_RE = re.compile(r"^\s*(?:savol\s*)?#?(\d+)\s*[\).\-:]\s*(.+)$", re.I)

KEY_HEADER_RE = re.compile(
    r"^\s*(javoblar(?:\s*kaliti)?|javob\s*kaliti|kalit(?:lar)?|to'?g'?ri\s*javoblar|answer\s*keys?|answers?|kalit\s*javoblar)\s*[:\-.]?\s*$",
    re.I,
)
KEY_PAIR_RE = re.compile(r"№?\s*(\d{1,3})\s*[-).:–—]?\s*([ABCDabcd])(?![A-Za-zА-я])")


def _is_complete(cur):
    o = cur["opts"]
    return bool(cur["text"]) and all(o.get(L) for L in ("A", "B", "C", "D")) and cur["correct"] in ("A", "B", "C", "D")


def _answer_key_pairs(line):
    pairs = {}
    for num, let in KEY_PAIR_RE.findall(line):
        pairs[int(num)] = let.upper()
    if not pairs:
        return None
    leftover = KEY_PAIR_RE.sub(" ", line)
    leftover = re.sub(r"(?i)(javoblar|javob|kalit|answer|key|to'?g'?ri|lar)", " ", leftover)
    leftover = re.sub(r"[\s,;.\-–—)|:()\[\]№#]+", "", leftover)
    return pairs if len(leftover) <= 2 else None


def _new_q(num=None, text=""):
    return {"num": num, "text": text, "opts": {}, "correct": ""}


def _try_split_inline(line):
    """Inline variantlarni ajratish: '1. Savol? A) x B) y C) z D) w Javob: B' -> alohida qatorlar."""
    # Avval oxiridagi javob qismini ajratish ("Javob: B", "Answer: C" kabi)
    answer_part = None
    m_ans = re.search(
        r"\s+((?:javob(?:i)?|to'?g'?ri\s*javob|answer|correct|kalit)\s*[:\-]?\s*[ABCDabcd])\s*$",
        line, re.I,
    )
    if m_ans:
        answer_part = m_ans.group(1).strip()
        line = line[:m_ans.start()]

    parts = re.split(r"(?<![A-Za-zА-я])([ABCDabcd])\s*[\).\-:\]]\s*", line)
    if len(parts) < 7:
        return None

    letters = [parts[i].upper() for i in range(1, len(parts) - 1, 2)]
    if len(letters) < 3:
        return None

    expected = list("ABCD")
    seq = []
    ei = 0
    for i, L in enumerate(letters):
        if ei < len(expected) and L == expected[ei]:
            seq.append(i)
            ei += 1
    if len(seq) < 3:
        return None

    result = []
    before = parts[0].strip()
    if before:
        result.append(before)

    for i in range(1, len(parts) - 1, 2):
        letter = parts[i]
        text = parts[i + 1].strip().rstrip(",;") if i + 1 < len(parts) else ""
        if text or letter.upper() in "ABCD":
            result.append(f"{letter.upper()}) {text}")

    if answer_part:
        result.append(answer_part)

    return result if len(result) >= 4 else None


def _preprocess_lines(lines):
    """Qatorlarni oldindan tayyorlash: inline variantlarni ajratish, OCR xatolarini tuzatish."""
    expanded = []
    for raw in lines:
        line = _clean(raw)
        if not line:
            expanded.append("")
            continue

        # Tab bilan ajratilgan ustunlarni alohida qatorlarga ajratish
        if "\t" in line and line.count("\t") >= 3:
            cells = [c.strip() for c in line.split("\t")]
            for c in cells:
                if c:
                    expanded.append(c)
            continue

        # Inline variantlarni ajratish (A) x B) y C) z D) w)
        split = _try_split_inline(line)
        if split:
            expanded.extend(split)
            continue

        # Javob qatorini matn bilan birlashgan holatda ajratish
        # Masalan: "D) to'rtinchi Javob: B" -> "D) to'rtinchi" va "Javob: B"
        m = re.match(r"(.+?)\s+((?:javob|to'?g'?ri|answer|correct|kalit)\s*[:\-]?\s*[ABCDabcd])\s*$", line, re.I)
        if m and OPT_RE.match(m.group(1)):
            expanded.append(m.group(1))
            expanded.append(m.group(2))
            continue

        expanded.append(line)
    return expanded


def parse_text_lines(lines):
    lines = _preprocess_lines(lines)

    questions = []
    cur = _new_q()
    answer_key = {}
    key_seq = []
    in_key_section = False

    def flush():
        if cur["text"]:
            questions.append(cur)

    for raw in lines:
        line = _clean(raw)
        if not line:
            continue

        if KEY_HEADER_RE.match(line):
            if cur["text"]:
                flush(); cur = _new_q()
            in_key_section = True
            continue

        kp = _answer_key_pairs(line)
        if kp is not None:
            if cur["text"]:
                flush(); cur = _new_q()
            answer_key.update(kp)
            in_key_section = True
            continue

        if in_key_section:
            bm = BARE_ANS_RE.match(line)
            if bm:
                key_seq.append(bm.group(1).upper())
                continue

        mark_correct = False
        m_lead = re.match(r"^\s*[\*★✓✔•●]+\s*([ABCDabcd])\b", line)
        if m_lead:
            mark_correct = True
            line = re.sub(r"^\s*[\*★✓✔•●]+\s*", "", line)

        am = ANS_RE.match(line)
        om = OPT_RE.match(line)
        qm = Q_RE.match(line)
        bare = BARE_ANS_RE.match(line)

        if am:
            cur["correct"] = am.group(2).upper()
        elif om:
            in_key_section = False
            L = om.group(1).upper()
            txt = om.group(2).strip()
            if re.search(r"[\*★✓✔•●]\s*$", txt):
                mark_correct = True
                txt = re.sub(r"\s*[\*★✓✔•●]+\s*$", "", txt).strip()
            # "+" belgisi ham to'g'ri javobni bildirishi mumkin
            if re.search(r"\s*\+\s*$", txt) and not cur["correct"]:
                mark_correct = True
                txt = re.sub(r"\s*\+\s*$", "", txt).strip()
            cur["opts"][L] = txt
            if mark_correct and not cur["correct"]:
                cur["correct"] = L
        elif qm:
            in_key_section = False
            if cur["text"]:
                flush()
            cur = _new_q(int(qm.group(1)) if qm.group(1).isdigit() else None, qm.group(2).strip())
        elif bare and cur["opts"] and not cur["correct"]:
            cur["correct"] = bare.group(1).upper()
        elif not cur["text"]:
            cur = _new_q(None, line)
        elif not cur["opts"]:
            cur["text"] += " " + line
        elif _is_complete(cur):
            flush()
            cur = _new_q(None, line)
        else:
            for L in ("D", "C", "B", "A"):
                if cur["opts"].get(L):
                    cur["opts"][L] += " " + line
                    break
            else:
                cur["text"] += " " + line
    if cur["text"]:
        flush()

    if key_seq and not answer_key:
        for i, L in enumerate(key_seq, start=1):
            answer_key[i] = L

    out = []
    for idx, q in enumerate(questions, start=1):
        correct = q["correct"]
        if correct not in ("A", "B", "C", "D"):
            kl = answer_key.get(q["num"]) or answer_key.get(idx)
            if kl in ("A", "B", "C", "D"):
                correct = kl
        o = q["opts"]
        built = _q(q["text"], o.get("A"), o.get("B"), o.get("C"), o.get("D"), correct)
        if built:
            out.append(built)
    return out


def parse_file(filename, data):
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        return parse_excel(data)
    if name.endswith(".csv"):
        return parse_csv(data)
    if name.endswith(".docx"):
        return parse_docx(data)
    if name.endswith(".pdf"):
        return parse_pdf(data)
    try:
        return parse_text_lines(data.decode("utf-8", errors="replace").splitlines())
    except Exception:
        return []
